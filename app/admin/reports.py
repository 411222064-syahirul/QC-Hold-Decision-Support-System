import csv
from io import StringIO
from flask import Blueprint, jsonify, request, render_template, make_response
from flask_login import login_required
from sqlalchemy import func, extract
from datetime import datetime

from app.models import db, QcHold, DecisionAction, DefectCategory
from app.auth import require_role

bp = Blueprint('admin_reports', __name__, url_prefix='/admin')

@bp.route('/reports')
@login_required
@require_role('admin', 'kepala_qc')
def reports_page():
    return render_template('admin/reports.html')

@bp.route('/api/reports/monthly', methods=['GET'])
@login_required
@require_role('admin', 'kepala_qc')
def get_monthly_report():
    month_str = request.args.get('month') # Format YYYY-MM
    if not month_str:
        return jsonify({'error': 'Parameter month diperlukan'}), 400
        
    year, month = month_str.split('-')
    
    # Base query for the month
    q = QcHold.query.filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    )
    
    # Basic Stats
    total_qc = q.count()
    total_ng = db.session.query(func.sum(QcHold.qty_ng)).filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).scalar() or 0
    total_ok = db.session.query(func.sum(QcHold.qty_ok)).filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).scalar() or 0
    
    # Decision distribution
    decision_dist_raw = db.session.query(
        DecisionAction.label, func.count(QcHold.id)
    ).join(QcHold, DecisionAction.id == QcHold.decision_action_id).filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).group_by(DecisionAction.label).all()
    
    decision_dist = [{'action': d[0], 'count': d[1]} for d in decision_dist_raw]
    
    # Defect category distribution
    defect_dist_raw = db.session.query(
        DefectCategory.name, func.count(QcHold.id)
    ).join(QcHold, DefectCategory.id == QcHold.defect_category_id).filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).group_by(DefectCategory.name).all()
    
    defect_dist = [{'category': d[0], 'count': d[1]} for d in defect_dist_raw]
    
    # Daily trend
    daily_trend_raw = db.session.query(
        func.date(QcHold.inspection_date), func.count(QcHold.id)
    ).filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).group_by(func.date(QcHold.inspection_date)).order_by(func.date(QcHold.inspection_date)).all()
    
    daily_trend = [{'date': str(d[0]), 'count': d[1]} for d in daily_trend_raw]
    
    # Ambil records untuk tabel laporan
    records = QcHold.query.filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).order_by(QcHold.inspection_date).all()
    
    records_data = []
    for r in records:
        records_data.append({
            'qc_code': r.qc_code,
            'inspection_date': str(r.inspection_date),
            'product_category': r.product_category.name if r.product_category else '-',
            'defect_category': r.defect_category.name if r.defect_category else '-',
            'qty_total': r.qty_total,
            'qty_ng': r.qty_ng,
            'decision': r.decision_action.code if r.decision_action else '-'
        })
    
    return jsonify({
        'total_qc_hold': total_qc,
        'total_ng': int(total_ng),
        'total_ok': int(total_ok),
        'decision_distribution': decision_dist,
        'defect_distribution': defect_dist,
        'daily_trend': daily_trend,
        'records': records_data
    })

@bp.route('/api/reports/export-excel', methods=['GET'])
@login_required
@require_role('admin', 'kepala_qc')
def export_excel():
    month_str = request.args.get('month') # Format YYYY-MM
    if not month_str:
        return jsonify({'error': 'Parameter month diperlukan'}), 400
        
    year, month = month_str.split('-')
    
    records = QcHold.query.filter(
        extract('year', QcHold.inspection_date) == year,
        extract('month', QcHold.inspection_date) == month,
        ~QcHold.qc_code.like('QC-IMP-%')
    ).order_by(QcHold.inspection_date).all()
    
    import pandas as pd
    import io
    
    data = []
    for r in records:
        data.append({
            'NO QC HOLD': r.qc_code,
            'TANGGAL': str(r.inspection_date),
            'PRODUK': r.product_category.name if r.product_category else '',
            'KATEGORI DEFECT': r.defect_category.name if r.defect_category else '',
            'LEVEL MASALAH': r.defect_severity,
            'QTY TOTAL': r.qty_total,
            'QTY NG': r.qty_ng,
            'KEPUTUSAN': r.decision_action.label if r.decision_action else '',
            'KETERANGAN AWAL': r.defect_description or '',
            'CATATAN PERBAIKAN': r.notes if r.status != 'OPEN' and r.notes else '',
            'PEMBUAT': r.creator.full_name if r.creator and r.creator.full_name else (r.creator.username if r.creator else ''),
            'TERAKHIR EDIT': r.updater.full_name if r.updater and r.updater.full_name else (r.updater.username if r.updater else ''),
            'PENENTU KEPUTUSAN': r.action_taker.full_name if r.action_taker and r.action_taker.full_name else (r.action_taker.username if r.action_taker else '')
        })
        
    df = pd.DataFrame(data)
    
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='QC_Hold_Data')
        
        # Styling the Excel worksheet
        worksheet = writer.sheets['QC_Hold_Data']
        
        try:
            from openpyxl.worksheet.table import Table, TableStyleInfo
            from openpyxl.utils import get_column_letter
            
            # Format as Excel Table
            num_rows = len(df) + 1
            num_cols = len(df.columns)
            if num_rows > 1:
                end_col_letter = get_column_letter(num_cols)
                table_ref = f"A1:{end_col_letter}{num_rows}"
                
                tab = Table(displayName="TabelQC", ref=table_ref)
                style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
                tab.tableStyleInfo = style
                worksheet.add_table(tab)
            
            # Auto-adjust column widths
            for col in worksheet.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2) * 1.2
                worksheet.column_dimensions[col_letter].width = min(adjusted_width, 50)  # max width 50
        except Exception as e:
            pass # fallback to plain excel if styling fails
        
    out.seek(0)
    output = make_response(out.getvalue())
    output.headers["Content-Disposition"] = f"attachment; filename=Laporan_QC_{month_str}.xlsx"
    output.headers["Content-type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return output


@bp.route('/api/reports/export-pdf', methods=['GET'])
@login_required
@require_role('admin', 'kepala_qc')
def export_pdf():
    from flask_login import current_user
    month_str = request.args.get('month')  # Format YYYY-MM
    if not month_str:
        return jsonify({'error': 'Parameter month diperlukan'}), 400

    year, month = month_str.split('-')

    # ── Query data ────────────────────────────────────────────────────────────
    base_filter = [
        extract('year', QcHold.inspection_date) == int(year),
        extract('month', QcHold.inspection_date) == int(month),
        ~QcHold.qc_code.like('QC-IMP-%'),
    ]

    records = QcHold.query.filter(*base_filter).order_by(QcHold.inspection_date).all()

    total_qc = len(records)
    total_ng = sum(r.qty_ng or 0 for r in records)
    total_ok = sum(r.qty_ok or 0 for r in records)

    decision_dist_raw = db.session.query(
        DecisionAction.label, func.count(QcHold.id)
    ).join(QcHold, DecisionAction.id == QcHold.decision_action_id).filter(
        *base_filter
    ).group_by(DecisionAction.label).all()

    # ── Build PDF ─────────────────────────────────────────────────────────────
    import io
    import os
    import locale
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm, cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                     Paragraph, Spacer, Image, HRFlowable)

    buf = io.BytesIO()
    page_w, page_h = landscape(A4)
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    # Custom styles
    s_company = ParagraphStyle('Company', parent=styles['Normal'],
                                fontSize=14, fontName='Helvetica-Bold',
                                alignment=TA_CENTER, leading=18)
    s_dept = ParagraphStyle('Dept', parent=styles['Normal'],
                             fontSize=10, fontName='Helvetica',
                             alignment=TA_CENTER, leading=14,
                             textColor=colors.HexColor('#444444'))
    s_title = ParagraphStyle('Title2', parent=styles['Normal'],
                              fontSize=13, fontName='Helvetica-Bold',
                              alignment=TA_CENTER, leading=18,
                              spaceBefore=6, spaceAfter=2)
    s_sub = ParagraphStyle('Sub', parent=styles['Normal'],
                            fontSize=9, fontName='Helvetica',
                            alignment=TA_CENTER, leading=12,
                            textColor=colors.HexColor('#555555'))
    s_label = ParagraphStyle('Label', parent=styles['Normal'],
                              fontSize=8, fontName='Helvetica',
                              textColor=colors.HexColor('#555555'), leading=11)
    s_value = ParagraphStyle('Value', parent=styles['Normal'],
                              fontSize=8, fontName='Helvetica-Bold', leading=11)
    s_cell = ParagraphStyle('Cell', parent=styles['Normal'],
                             fontSize=7, fontName='Helvetica', leading=10,
                             wordWrap='CJK')
    s_cell_bold = ParagraphStyle('CellB', parent=s_cell,
                                  fontName='Helvetica-Bold')
    s_footer = ParagraphStyle('Footer', parent=styles['Normal'],
                               fontSize=7, fontName='Helvetica',
                               textColor=colors.HexColor('#888888'),
                               alignment=TA_CENTER)

    elements = []

    # ── HEADER: Logo + Company + Dept ─────────────────────────────────────────
    logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                             'static', 'images', 'logo.png')

    header_data = []
    if os.path.exists(logo_path):
        logo = Image(logo_path, width=1.8 * cm, height=1.8 * cm)
        company_block = [
            Paragraph('PT. SS HEP', s_company),
            Paragraph('Departemen Quality Control', s_dept),
        ]
        header_data = [[logo, company_block]]
    else:
        company_block = [
            Paragraph('PT. SS HEP', s_company),
            Paragraph('Departemen Quality Control', s_dept),
        ]
        header_data = [['', company_block]]

    avail_w = page_w - 3 * cm  # total available width
    header_table = Table(header_data, colWidths=[2.5 * cm, avail_w - 2.5 * cm])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (0, 0), 'CENTER'),
        ('ALIGN', (1, 0), (1, 0), 'CENTER'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(header_table)

    # Divider line
    elements.append(Spacer(1, 3 * mm))
    elements.append(HRFlowable(width='100%', thickness=1.5,
                                color=colors.HexColor('#333333'),
                                spaceAfter=2 * mm))

    # ── TITLE + PERIOD ────────────────────────────────────────────────────────
    # Build Indonesian month name
    bulan_map = {
        1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April',
        5: 'Mei', 6: 'Juni', 7: 'Juli', 8: 'Agustus',
        9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'
    }
    bulan_nama = bulan_map.get(int(month), month)
    periode_text = f'{bulan_nama} {year}'

    elements.append(Paragraph('LAPORAN HASIL KEPUTUSAN QC HOLD', s_title))
    elements.append(Paragraph(f'Periode: {periode_text}', s_sub))
    elements.append(Spacer(1, 4 * mm))

    # ── INFO PENCETAK ─────────────────────────────────────────────────────────
    now_str = datetime.now().strftime('%d %B %Y, %H:%M WIB').replace(
        datetime.now().strftime('%B'),
        bulan_map.get(datetime.now().month, datetime.now().strftime('%B'))
    )
    user_nip = current_user.nip or '-'
    user_name = current_user.full_name or current_user.username
    user_username = current_user.username

    info_data = [
        [Paragraph('Tanggal Cetak', s_label), Paragraph(f': {now_str}', s_value),
         Paragraph('NIP', s_label), Paragraph(f': {user_nip}', s_value)],
        [Paragraph('Dicetak Oleh', s_label), Paragraph(f': {user_name}', s_value),
         Paragraph('Username', s_label), Paragraph(f': {user_username}', s_value)],
    ]
    info_table = Table(info_data, colWidths=[2.5 * cm, 6 * cm, 2 * cm, 6 * cm])
    info_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 1),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 5 * mm))

    # ── SUMMARY TABLE ─────────────────────────────────────────────────────────
    elements.append(Paragraph('<b>Ringkasan</b>', ParagraphStyle(
        'SH', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold',
        spaceAfter=2 * mm)))

    summary_data = [
        [Paragraph('<b>Total QC Hold</b>', s_cell_bold),
         Paragraph('<b>Total NG</b>', s_cell_bold),
         Paragraph('<b>Total OK</b>', s_cell_bold)]
    ]
    # Add decision columns
    for label, count in decision_dist_raw:
        summary_data[0].append(Paragraph(f'<b>{label}</b>', s_cell_bold))

    summary_row = [
        Paragraph(str(total_qc), s_cell),
        Paragraph(str(total_ng), s_cell),
        Paragraph(str(total_ok), s_cell),
    ]
    for label, count in decision_dist_raw:
        summary_row.append(Paragraph(str(count), s_cell))
    summary_data.append(summary_row)

    n_sum_cols = len(summary_data[0])
    sum_col_w = avail_w / n_sum_cols
    summary_table = Table(summary_data, colWidths=[sum_col_w] * n_sum_cols)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a56db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f6f8fa')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d0d7de')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 6 * mm))

    # ── DETAIL DATA TABLE ─────────────────────────────────────────────────────
    elements.append(Paragraph('<b>Detail Data QC Hold</b>', ParagraphStyle(
        'DH', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold',
        spaceAfter=2 * mm)))

    col_headers = ['No', 'No QC Hold', 'Tanggal', 'Produk', 'Part No',
                   'Jenis Masalah', 'Severity', 'Qty Total', 'Qty NG',
                   'Keputusan', 'Status', 'Keterangan']
    table_data = [[Paragraph(f'<b>{h}</b>', s_cell_bold) for h in col_headers]]

    severity_map = {1: 'Low', 2: 'Medium', 3: 'High'}
    for idx, r in enumerate(records, 1):
        insp_date_str = r.inspection_date.strftime('%d-%m-%Y') if r.inspection_date else '-'
        table_data.append([
            Paragraph(str(idx), s_cell),
            Paragraph(r.qc_code or '-', s_cell),
            Paragraph(insp_date_str, s_cell),
            Paragraph(r.product_category.name if r.product_category else '-', s_cell),
            Paragraph(r.part_no or '-', s_cell),
            Paragraph(r.defect_category.name if r.defect_category else '-', s_cell),
            Paragraph(severity_map.get(r.defect_severity, str(r.defect_severity or '-')), s_cell),
            Paragraph(str(r.qty_total or 0), s_cell),
            Paragraph(str(r.qty_ng or 0), s_cell),
            Paragraph(r.decision_action.label if r.decision_action else '-', s_cell),
            Paragraph(r.status or '-', s_cell),
            Paragraph(r.defect_description or '-', s_cell),
        ])

    # Column widths proportional
    detail_widths = [
        0.6 * cm,   # No
        2.2 * cm,   # No QC Hold
        1.8 * cm,   # Tanggal
        2.2 * cm,   # Produk
        1.8 * cm,   # Part No
        2.5 * cm,   # Jenis Masalah
        1.3 * cm,   # Severity
        1.3 * cm,   # Qty Total
        1.2 * cm,   # Qty NG
        1.8 * cm,   # Keputusan
        1.2 * cm,   # Status
    ]
    used = sum(detail_widths)
    detail_widths.append(avail_w - used)  # Keterangan gets remaining space

    detail_table = Table(table_data, colWidths=detail_widths, repeatRows=1)
    detail_table.setStyle(TableStyle([
        # Header row
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a56db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        # Alternating row colors
        *[('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f6f8fa'))
          for i in range(2, len(table_data), 2)],
        # Grid & padding
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#d0d7de')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),      # No column centered
        ('ALIGN', (7, 1), (8, -1), 'CENTER'),       # Qty columns centered
        ('ALIGN', (10, 1), (10, -1), 'CENTER'),     # Status centered
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(detail_table)

    # ── FOOTER NOTE ───────────────────────────────────────────────────────────
    elements.append(Spacer(1, 8 * mm))
    elements.append(HRFlowable(width='100%', thickness=0.5,
                                color=colors.HexColor('#cccccc'),
                                spaceAfter=2 * mm))
    elements.append(Paragraph(
        f'Dokumen ini dicetak secara otomatis oleh Sistem QC Hold DSS — '
        f'PT. SS HEP © {year}',
        s_footer
    ))

    # ── Generate ──────────────────────────────────────────────────────────────
    doc.build(elements)
    buf.seek(0)

    response = make_response(buf.getvalue())
    response.headers['Content-Disposition'] = (
        f'attachment; filename=Laporan_QC_Hold_{month_str}.pdf'
    )
    response.headers['Content-Type'] = 'application/pdf'
    return response
